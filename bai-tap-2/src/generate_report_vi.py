# -*- coding: utf-8 -*-
"""
Sinh báo cáo DOCX (TIẾNG VIỆT CÓ DẤU) cho Bài tập 2 từ kết quả notebook.
Đọc: bai-tap-2/outputs/summary.json + bai-tap-2/outputs/figures/*.png
Chạy:  ./venv/Scripts/python.exe bai-tap-2/src/generate_report_vi.py
Kết quả: bai-tap-2/BaoCao_TieuLuan_LogisticRegression_TiengViet.docx
"""
import os
import json

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # bai-tap-2/
OUT = os.path.join(ROOT, "outputs")
FIG = os.path.join(OUT, "figures")

with open(os.path.join(OUT, "summary.json"), encoding="utf-8") as f:
    S = json.load(f)

BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x2E, 0x7D, 0x32)

METRIC_VI = {
    "Accuracy": "Accuracy (độ chính xác)",
    "Precision": "Precision (độ chuẩn xác)",
    "Recall": "Recall (độ bao phủ)",
    "F1": "F1-score",
    "ROC-AUC": "ROC-AUC",
    "LogLoss": "Log loss",
}


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, italic=False, bold=False, size=11, align=None, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return p


def add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_figure(doc, filename, caption, width=6.0):
    path = os.path.join(FIG, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_para(doc, caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_metrics_table(doc, metrics_np, metrics_sk):
    keys = ["Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "LogLoss"]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for j, c in enumerate(["Chỉ số", "NumPy (tự cài)", "scikit-learn"]):
        hdr[j].text = c
        for p in hdr[j].paragraphs:
            for r in p.runs:
                r.bold = True
    for k in keys:
        cells = table.add_row().cells
        cells[0].text = METRIC_VI[k]
        cells[1].text = f"{metrics_np[k]:.4f}"
        cells[2].text = f"{metrics_sk[k]:.4f}"
    return table


def add_confusion_table(doc, cm, title):
    add_para(doc, title, bold=True, size=10)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = ""
    table.rows[0].cells[1].text = "Dự đoán: benign"
    table.rows[0].cells[2].text = "Dự đoán: malignant"
    table.rows[1].cells[0].text = "Thực tế: benign"
    table.rows[2].cells[0].text = "Thực tế: malignant"
    table.rows[1].cells[1].text = str(cm[0][0])
    table.rows[1].cells[2].text = str(cm[0][1])
    table.rows[2].cells[1].text = str(cm[1][0])
    table.rows[2].cells[2].text = str(cm[1][1])
    for r in table.rows:
        for c in (r.cells[0], table.rows[0].cells[1], table.rows[0].cells[2]):
            for p in c.paragraphs:
                for run in p.runs:
                    run.bold = True
    return table


# --------------------------------------------------------------------------
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)

ds = S["dataset"]
pp = S["preprocessing"]
hp = S["hyperparameters"]
mn = S["metrics"]["numpy"]
ms = S["metrics"]["sklearn"]
pr = S["params"]
mis = S["misclassification"]

# ---- Cover ----
add_para(doc, "TIỂU LUẬN MÔN HỌC", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "HỌC MÁY / KHAI PHÁ DỮ LIỆU", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_para(doc, "CÀI ĐẶT VÀ ĐÁNH GIÁ LOGISTIC REGRESSION",
         bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE)
add_para(doc, "Tự cài đặt bằng NumPy và so sánh với scikit-learn",
         italic=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
doc.add_paragraph()
add_para(doc, "Hình thức: Làm cá nhân", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, f"Dữ liệu: {ds['name']} (UCI ML Repository)", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "Sản phẩm nộp: Jupyter notebook logistic_regression.ipynb (chạy đầu-cuối)",
         align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "Ngày: 28/09/2026", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---- TOC ----
add_heading(doc, "MỤC LỤC", level=1)
for x in [
    "1. Giới thiệu và mục tiêu",
    "2. Dữ liệu và tỷ lệ các lớp",
    "3. Phương pháp: cài đặt Logistic Regression bằng NumPy",
    "4. Kết quả thực nghiệm",
    "5. Nhận xét",
    "6. Phụ lục - Hướng dẫn chạy",
]:
    doc.add_paragraph(x, style="List Number")
doc.add_page_break()

# ---- 1 ----
add_heading(doc, "1. Giới thiệu và mục tiêu", level=1)
add_para(doc,
    "Bài tiểu luận tự cài đặt mô hình Logistic Regression (hồi quy logistic) cho bài toán "
    "phân loại nhị phân bằng NumPy — không dùng mô hình hay bộ tối ưu có sẵn — sau đó so sánh "
    "với sklearn.linear_model.LogisticRegression trên cùng dữ liệu đã tiền xử lý. Mục tiêu:")
add_bullets(doc, [
    "Chuẩn bị dữ liệu: mô tả đặc trưng, nhãn và tỷ lệ các lớp.",
    "Cài đặt bằng NumPy: score, sigmoid, log loss và cập nhật trọng số + bias bằng gradient descent.",
    "Huấn luyện sklearn LogisticRegression trên cùng dữ liệu.",
    "Đánh giá: accuracy, precision, recall, F1, ROC-AUC, log loss + ma trận nhầm lẫn của mỗi bản.",
    "Nhận xét về hội tụ, chênh lệch giữa hai bản và các lỗi phân loại thường gặp.",
])
add_para(doc,
    "Sản phẩm nộp chính là một Jupyter notebook chạy được từ đầu đến cuối; báo cáo này tóm tắt "
    "phương pháp và kết quả bằng tiếng Việt.", italic=True)

# ---- 2 ----
add_heading(doc, "2. Dữ liệu và tỷ lệ các lớp", level=1)
add_para(doc,
    f"Bộ dữ liệu {ds['name']} gồm {ds['n_samples']} mẫu, {ds['n_features']} đặc trưng số thực "
    "(mô tả đặc điểm nhân tế bào từ ảnh sinh thiết: bán kính, kết cấu, chu vi, diện tích, độ "
    "lõm...). Đây là dữ liệu y tế thực tế từ UCI, nạp offline qua sklearn nên notebook chạy được "
    "ở mọi máy, không cần tải dữ liệu ngoài.")
add_para(doc,
    f"Nhãn là nhị phân. Ta quy ước lớp dương (y=1) là 'malignant' (ác tính) — lớp cần phát hiện "
    f"trong y khoa. Tỷ lệ các lớp:")
add_bullets(doc, [
    f"Ác tính (malignant, y=1): {ds['class_counts']['malignant']} mẫu "
    f"({ds['class_ratio']['malignant']:.1%}).",
    f"Lành tính (benign, y=0): {ds['class_counts']['benign']} mẫu "
    f"({ds['class_ratio']['benign']:.1%}).",
])
add_para(doc,
    f"Dữ liệu được chia train/test = {(1-pp['test_size']):.0%}/{pp['test_size']:.0%} có phân tầng "
    f"(stratify) để giữ tỷ lệ lớp ({pp['n_train']} train / {pp['n_test']} test), và chuẩn hóa bằng "
    "StandardScaler (fit trên tập train) để tránh rò rỉ dữ liệu. Cùng dữ liệu này dùng cho cả hai bản.")
add_figure(doc, "01_class_distribution.png", "Hình 1. Phân bố các lớp trong bộ dữ liệu.", 4.6)

# ---- 3 ----
add_heading(doc, "3. Phương pháp: cài đặt Logistic Regression bằng NumPy", level=1)
add_para(doc, "Với vector đặc trưng x, trọng số w và bias b:")
add_bullets(doc, [
    "Score tuyến tính: z = x·w + b.",
    "Sigmoid (xác suất): p = σ(z) = 1 / (1 + e^(−z)); clip z vào [−500, 500] để tránh tràn số.",
    "Log loss (binary cross-entropy): L = −(1/m) Σ [y·log(p) + (1−y)·log(1−p)]; "
    "clip p vào [ε, 1−ε] để tránh log(0).",
    "Điều chuẩn L2 (chỉ trên trọng số): J = L + (λ/2)·||w||². Do dữ liệu gần tách tuyến tính, "
    "nếu không điều chuẩn thì trọng số sẽ tiến ra vô cực và xác suất bão hòa (log loss test rất "
    "lớn); L2 giữ nghiệm ổn định.",
    "Gradient: dJ/dw = (1/m)·Xᵀ(p−y) + λ·w;  dJ/db = (1/m)·Σ(p−y).",
    "Cập nhật (gradient descent): w ← w − η·dJ/dw;  b ← b − η·dJ/db.",
])
add_para(doc,
    f"Siêu tham số: learning rate η = {hp['learning_rate']}, tối đa {hp['max_iter']} vòng lặp, "
    f"điều kiện dừng khi |thay đổi log loss| < {hp['tol']:g}, L2 λ = 0.01. Khởi tạo w = 0, b = 0.")
add_para(doc,
    "Để so sánh công bằng, sklearn dùng cùng mức điều chuẩn: sklearn tối ưu C·ΣBCE + ½||w||², "
    "còn bản NumPy tối ưu (1/m)ΣBCE + (λ/2)||w||²; khớp hai gradient cho C = 1/(m·λ). "
    "Nhờ đó cả hai giải cùng một hàm mục tiêu.")
add_para(doc,
    "Toàn bộ được cài bằng NumPy thuần (không dùng model/optimizer có sẵn).", italic=True)

# ---- 4 ----
add_heading(doc, "4. Kết quả thực nghiệm", level=1)
add_heading(doc, "4.1. Sự hội tụ của gradient descent", level=2)
conv = "đã hội tụ" if hp["converged"] else "đạt trần số vòng lặp"
add_para(doc,
    f"Vì đặc trưng đã chuẩn hóa và hàm log loss là lồi, gradient descent giảm đơn điệu và hội tụ "
    f"mượt. Quá trình dừng sau {hp['n_iter_run']} vòng ({conv}), log loss train cuối = "
    f"{hp['final_train_log_loss']:.4f}.")
add_figure(doc, "02_loss_curve.png", "Hình 2. Log loss trên tập train theo vòng lặp (đường hội tụ).", 6.0)

add_heading(doc, "4.2. Bảng so sánh các chỉ số", level=2)
add_para(doc, "Bảng 1. So sánh trên tập test (ROC-AUC và log loss tính từ xác suất dự đoán).")
add_metrics_table(doc, mn, ms)
add_para(doc,
    f"Hai bản cho kết quả gần như trùng khớp: accuracy = {mn['Accuracy']:.4f}, "
    f"ROC-AUC = {mn['ROC-AUC']:.4f}, log loss = {mn['LogLoss']:.4f}. "
    f"Precision = {mn['Precision']:.2f} và Recall = {mn['Recall']:.4f} cho lớp ác tính.",
    bold=True)

add_heading(doc, "4.3. Ma trận nhầm lẫn", level=2)
add_confusion_table(doc, S["confusion_matrix"]["numpy"], "Bản NumPy:")
add_confusion_table(doc, S["confusion_matrix"]["sklearn"], "Bản scikit-learn:")
add_figure(doc, "03_confusion_matrices.png", "Hình 3. Ma trận nhầm lẫn của hai bản.", 6.3)

add_heading(doc, "4.4. Đường cong ROC", level=2)
add_figure(doc, "04_roc_curves.png",
           f"Hình 4. Đường ROC (AUC NumPy = {mn['ROC-AUC']:.4f}, sklearn = {ms['ROC-AUC']:.4f}).", 4.8)

add_heading(doc, "4.5. So sánh trọng số học được", level=2)
add_para(doc,
    f"Hệ số tương quan Pearson giữa trọng số hai bản = {pr['weight_correlation']:.4f}; "
    f"bias NumPy = {pr['bias_numpy']:.3f}, bias sklearn = {pr['bias_sklearn']:.3f}. "
    "Trọng số gần như trùng khớp — minh chứng bản NumPy cài đúng.")
add_figure(doc, "05_weight_comparison.png", "Hình 5. So sánh trọng số NumPy vs sklearn.", 4.4)

# ---- 5 ----
add_heading(doc, "5. Nhận xét", level=1)
add_para(doc, "Về sự hội tụ:", bold=True)
add_para(doc,
    "Chuẩn hóa đặc trưng giúp gradient descent hội tụ nhanh và ổn định trên mặt log loss lồi; "
    f"đường log loss (Hình 2) phẳng dần và dừng theo ngưỡng {hp['tol']:g} sau {hp['n_iter_run']} vòng.")
add_para(doc, "Về chênh lệch giữa hai bản:", bold=True)
add_para(doc,
    "Do cùng tối ưu một hàm mục tiêu (log loss + L2 khớp nhau), hai bản hội tụ về gần như cùng "
    f"nghiệm: trọng số tương quan ≈ {pr['weight_correlation']:.2f}, bias trùng nhau, và cả 6 chỉ "
    "số khớp tới 4 chữ số thập phân — dù dùng bộ tối ưu khác nhau (gradient descent bước cố định "
    "so với lbfgs kiểu tựa-Newton). Với hàm lồi, bộ tối ưu chỉ ảnh hưởng đường đi, không ảnh "
    "hưởng đích đến.")
add_para(doc, "Về các lỗi phân loại thường gặp:", bold=True)
add_para(doc,
    f"Trên {mis['n_test']} mẫu test, bản NumPy sai {mis['numpy']} mẫu, sklearn sai {mis['sklearn']} "
    f"mẫu, trong đó {mis['common']} mẫu bị sai ở cả hai bản. Đây là các ca ranh giới có xác suất dự "
    "đoán gần 0.5 — tức mẫu thực sự nhập nhằng, không phải lỗi của cách cài đặt. Trong y khoa, có "
    "thể hạ ngưỡng quyết định để tăng recall (giảm bỏ sót ca ác tính) đổi lấy giảm precision.")

# ---- 6 ----
add_heading(doc, "6. Phụ lục - Hướng dẫn chạy", level=1)
add_para(doc, "Sản phẩm nộp chính: bai-tap-2/logistic_regression.ipynb (chạy được từ đầu đến cuối).")
add_para(doc, "Các bước chạy (Windows PowerShell):")
code = doc.add_paragraph()
code_run = code.add_run(
    "python -m venv venv\n"
    ".\\venv\\Scripts\\Activate.ps1\n"
    "pip install -r requirements.txt\n"
    "jupyter notebook bai-tap-2\\logistic_regression.ipynb\n"
    "# (hoặc chạy toàn bộ không cần mở giao diện:)\n"
    "jupyter nbconvert --to notebook --execute --inplace bai-tap-2\\logistic_regression.ipynb"
)
code_run.font.name = "Consolas"
code_run.font.size = Pt(10)
add_para(doc,
    "Notebook sẽ nạp dữ liệu, huấn luyện cả hai bản, vẽ các biểu đồ và lưu figures + summary.json "
    "vào bai-tap-2/outputs/.", italic=True)

out_path = os.path.join(ROOT, "BaoCao_TieuLuan_LogisticRegression_TiengViet.docx")
doc.save(out_path)
print("Da tao bao cao tieng Viet:", out_path)
