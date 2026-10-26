"""Sinh Slide_NhanDangMau.pptx (python-pptx) từ results/ và figures/."""
import json
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from members import MEMBERS, GROUP, SCHOOL, ADVISOR

hexc = RGBColor.from_string
WINE, INK, ROSE, GREY, WHITE, BLUE, SOFT = hexc("5A1E2B"), hexc("2B2024"), hexc("F6EDEF"), hexc("6B5F63"), hexc("FFFFFF"), hexc("3A7CA5"), hexc("E8C9CF")
R = "results/"
cmp_ = pd.read_csv(R + "model_comparison.csv"); S = json.load(open(R + "summary.json")); ded = pd.read_csv(R + "dedup_check.csv")
thr_all = pd.read_csv(R + "threshold_all_models.csv"); cms = json.load(open(R + "confusion_matrices.json"))
row = lambda n: cmp_[cmp_["Mô hình"] == n].iloc[0]
BASE, SVM = "Cơ sở: k-NN (k=5, không chuẩn hóa)", "SVM RBF (cân bằng lớp, tuned)"
b, s = row(BASE), row(SVM)
n3 = lambda v: f"{v:.3f}".replace(".", ",")
short = lambda n: (n.replace("Cơ sở: ", "[CS] ").replace(" (k=5, không chuẩn hóa)", " thô").replace("Tham chiếu: luôn đoán lớp đa số", "[TC] Lớp đa số")
                   .replace(" (cân bằng lớp, tuned)", "").replace(" (cân bằng lớp)", "").replace(" (tuned)", "").replace("Hồi quy logistic", "Logistic").replace("Chọn đặc trưng (ANOVA, k tuned) + SVM", "Chọn ĐT + SVM").replace("Hồi quy thứ bậc: SVR điểm → ngưỡng", "SVR thứ bậc"))

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5); blank = prs.slide_layouts[6]
def bg(sl, c): sl.background.fill.solid(); sl.background.fill.fore_color.rgb = c
def tb(sl, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    t = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); f = t.text_frame; f.word_wrap = True; f.vertical_anchor = anchor
    for i, l in enumerate(text if isinstance(text, list) else [text]):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph(); p.alignment = align; p.space_after = Pt(6)
        r = p.add_run(); r.text = l; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Calibri"
    return t
def bullets(sl, x, y, w, h, items, size=18, color=INK):
    t = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); f = t.text_frame; f.word_wrap = True
    for i, it in enumerate(items):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph(); p.space_after = Pt(10)
        r = p.add_run(); r.text = "▪  " + it; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
def card(sl, x, y, w, h, fill=ROSE):
    c = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); c.adjustments[0] = .06
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False; return c
def slide(title, notes=""):
    sl = prs.slides.add_slide(blank); bg(sl, WHITE)
    tb(sl, .6, .35, 12.1, .9, title, 32, True, WINE, anchor=MSO_ANCHOR.MIDDLE); sl.notes_slide.notes_text_frame.text = notes; return sl
def stat(sl, x, y, w, big, small, col=WINE, h=1.9):
    card(sl, x, y, w, h); tb(sl, x, y + .2, w, 1, big, 36, True, col, PP_ALIGN.CENTER); tb(sl, x + .15, y + 1.15, w - .3, h - 1.2, small, 14, False, GREY, PP_ALIGN.CENTER)
def pic(sl, f, x, y, w=None, h=None): sl.shapes.add_picture("figures/" + f, Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)
def table(sl, rows, x, y, widths, size=14):
    t = sl.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(sum(widths)), Inches(.42 * len(rows))).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v); p = c.text_frame.paragraphs[0]; p.font.size = Pt(size); p.font.bold = i == 0; p.font.name = "Calibri"
            c.fill.solid(); c.fill.fore_color.rgb = WINE if i == 0 else (ROSE if i % 2 else WHITE); p.font.color.rgb = WHITE if i == 0 else INK
    for j, w in enumerate(widths): t.columns[j].width = Inches(w)

# 1 Bìa
sl = prs.slides.add_slide(blank); bg(sl, WINE)
tb(sl, .8, 1.7, 11.7, 1.6, "Nhận dạng rượu vang chất lượng cao", 46, True, WHITE)
tb(sl, .8, 3.3, 11.7, 1, "Từ 11 chỉ tiêu hóa lý – bộ dữ liệu Wine Quality (UCI): phương án cơ sở so với cải tiến", 22, False, SOFT)
tb(sl, .8, 5.6, 11.7, 1.2, ["Bài tập nhóm cuối kỳ – Học phần Nhận dạng mẫu – Học viện Công nghệ Bưu chính Viễn thông", GROUP, f"Giảng viên hướng dẫn: {ADVISOR}"], 16, False, SOFT)
sl.notes_slide.notes_text_frame.text = "Giới thiệu nhóm và đề tài."

# 2 Bài toán
sl = slide("Bài toán: thay hội đồng nếm bằng phép đo hóa lý?", "Chuyển điểm 0–10 thành nhị phân: ≥7 là chất lượng cao.")
bullets(sl, .6, 1.6, 6.8, 4.8, ["Đầu vào: 11 chỉ tiêu hóa lý (độ cồn, axit, đường dư, SO₂, tỉ trọng…) + loại rượu đỏ/trắng",
                                "Đầu ra: chất lượng cao (điểm cảm quan ≥ 7) hay thường",
                                "Dữ liệu thật: rượu Vinho Verde, Bồ Đào Nha, 2004–2007 (Cortez et al., UCI #186)",
                                "Mục tiêu: vượt phương án cơ sở có kiểm định, hiểu mô hình sai ở đâu"], 19)
stat(sl, 7.8, 1.7, 2.4, f"{S['n_raw']:,}".replace(",", "."), "mẫu gốc"); stat(sl, 10.4, 1.7, 2.4, f"{S['n_dup']:,}".replace(",", "."), "bản ghi trùng → loại", BLUE)
stat(sl, 7.8, 3.9, 2.4, f"{S['n']:,}".replace(",", "."), "mẫu sử dụng"); stat(sl, 10.4, 3.9, 2.4, f"{S['pos']/S['n']:.0%}", "lớp chất lượng cao")

# 3 Dữ liệu
sl = slide("Dữ liệu: lớp thiểu số, chồng lấp mạnh, độ cồn dẫn dắt", "Phân bố điểm lệch, 6 và 7 nằm hai bên ranh giới.")
pic(sl, "quality_dist.png", .6, 1.5, h=3.7); pic(sl, "importance.png", 6.9, 1.5, h=3.7)
bullets(sl, .6, 5.4, 12.1, 1.8, ["Chỉ 19% mẫu đạt ≥ 7 → accuracy gây hiểu nhầm; dùng F1, recall, AUC",
                                 "Độ cồn quan trọng nhất; hai lớp chồng lấp trên mọi chỉ tiêu, PCA cần 9/12 thành phần cho 95% phương sai"], 17)

# 4 Phương án
sl = slide("Phương án cơ sở và cải tiến", "Mọi tiền xử lý nằm trong Pipeline.")
card(sl, .6, 1.6, 5.3, 5.0, hexc("E3EEF5")); tb(sl, .9, 1.75, 4.8, .6, "Cơ sở", 24, True, BLUE)
bullets(sl, .9, 2.5, 4.8, 3.9, ["k-NN (k=5) trên đặc trưng thô", "Naive Bayes Gaussian", "Tham chiếu: luôn đoán lớp đa số"], 19)
card(sl, 6.2, 1.6, 6.5, 5.0); tb(sl, 6.5, 1.75, 6, .6, "Cải tiến / thay thế", 24, True, WINE)
bullets(sl, 6.5, 2.5, 6, 4.0, ["Chuẩn hóa z-score", "Trọng số cân bằng lớp", "Logistic, LDA/QDA, SVM RBF, PCA + SVM", "Random Forest, Gradient Boosting", "Chọn đặc trưng; hồi quy thứ bậc (SVR)", "Chọn ngưỡng quyết định tối ưu F1"], 18)

# 5 Quy trình
sl = slide("Quy trình đánh giá", "Nested CV; Nadeau–Bengio vì các fold chồng lấp.")
for i, (a, c) in enumerate([("Loại trùng", f"{S['n_dup']} bản ghi – tránh rò rỉ"), ("5×3 CV", "phân tầng, 15 fold"), ("Nested", "GridSearch trong mỗi fold"), ("t-test", "hiệu chỉnh Nadeau–Bengio")]):
    stat(sl, .6 + i * 3.1, 1.8, 2.9, a, c, WINE if i else BLUE)
bullets(sl, .6, 4.2, 12.1, 2.6, ["Chỉ số chính: F1 và recall của lớp chất lượng cao, ROC-AUC; accuracy chỉ để tham khảo",
                                 "Phân tích lỗi trên dự đoán ngoài-fold cho toàn bộ 5.320 mẫu; seed = 42"], 19)

# 6 Kết quả chart
sl = slide("Kết quả: F1 tăng gần gấp đôi so với cơ sở", "Ba mô hình đầu chênh nhau nhỏ hơn độ lệch chuẩn fold.")
cd = CategoryChartData(); order = cmp_[~cmp_["Mô hình"].str.startswith("Tham chiếu")].sort_values("f1")
cd.categories = [short(n) for n in order["Mô hình"]]; cd.add_series("F1", [round(v, 3) for v in order["f1"]])
ch = sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(.6), Inches(1.4), Inches(7.6), Inches(5.6), cd).chart
ch.has_legend = False; pl = ch.plots[0]; pl.gap_width = 45; pl.has_data_labels = True
pl.data_labels.font.size = Pt(13); pl.data_labels.number_format = "0.000"; pl.data_labels.number_format_is_linked = False; pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
va = ch.value_axis; va.minimum_scale = 0; va.maximum_scale = .7; va.has_major_gridlines = False; va.tick_labels.font.size = Pt(12); ch.category_axis.tick_labels.font.size = Pt(13)
for i, n in enumerate(order["Mô hình"]):
    pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = BLUE if n.startswith("Cơ sở") else WINE
stat(sl, 8.7, 1.6, 4, f"{n3(b['f1'])} → {n3(s['f1'])}", "F1: k-NN thô → SVM RBF")
odr0 = row("Hồi quy thứ bậc: SVR điểm → ngưỡng")
stat(sl, 8.7, 3.7, 4, n3(odr0['f1']), "F1 cao nhất: SVR thứ bậc; mọi cải tiến p < 0,001")
tb(sl, 8.7, 5.8, 4, .8, "Xanh: cơ sở; đỏ rượu: cải tiến. F1 ở ngưỡng mặc định.", 13, False, GREY)

# 7 Accuracy bẫy
dm = row("Tham chiếu: luôn đoán lớp đa số")
sl = slide("Vì sao không dùng accuracy", "Accuracy của lớp đa số cao hơn SVM.")
table(sl, [["Mô hình", "Accuracy", "Recall", "F1"]] + [[short(n), f"{row(n)['acc']:.1%}".replace(".", ","), f"{row(n)['rec']:.1%}".replace(".", ","), n3(row(n)["f1"])] for n in ["Tham chiếu: luôn đoán lớp đa số", BASE, "k-NN + chuẩn hóa", SVM]], .6, 1.7, [3.6, 1.8, 1.6, 1.4], 16)
card(sl, 9.3, 1.7, 3.4, 3.6); tb(sl, 9.5, 1.9, 3.0, 3.3, ["Đoán 'thường' cho mọi chai:", f"accuracy {dm['acc']:.0%}", "nhưng không tìm được chai ngon nào"], 18, False, INK)
bullets(sl, .6, 4.6, 8.4, 2, ["Chuẩn hóa: F1 0,277 → 0,452 (SO₂ ~ 10² lấn át độ cồn ~ 10)", "Thêm cân bằng lớp (SVM): recall 36% → 81%"], 18)

# 7b Gợi ý của UCI
fsr, odr = row("Chọn đặc trưng (ANOVA, k tuned) + SVM"), row("Hồi quy thứ bậc: SVR điểm → ngưỡng")
sl = slide("Thử theo gợi ý của UCI: chọn đặc trưng, nhãn có thứ tự", "UCI: 'không chắc mọi biến đều liên quan'; lớp có thứ tự và không cân bằng.")
card(sl, .6, 1.6, 5.9, 4.9); tb(sl, .85, 1.75, 5.4, .6, "Chọn đặc trưng (ANOVA F + SVM)", 20, True, WINE)
tb(sl, .85, 2.45, 5.4, .9, f"k = {S['fs_k']}/12", 36, True, WINE)
bullets(sl, .85, 3.5, 5.4, 3, [f"F1 {n3(fsr['f1'])} so với SVM đầy đủ {n3(s['f1'])}", "Nested CV giữ lại gần như mọi chỉ tiêu: mỗi chỉ tiêu đóng góp một ít"], 16)
card(sl, 6.8, 1.6, 5.9, 4.9); tb(sl, 7.05, 1.75, 5.4, .6, "Hồi quy thứ bậc (SVR điểm → ngưỡng)", 20, True, WINE)
tb(sl, 7.05, 2.45, 5.4, .9, f"F1 {n3(odr['f1'])}", 36, True, WINE)
bullets(sl, 7.05, 3.5, 5.4, 3, [f"Recall {odr['rec']:.0%}, precision {odr['prec']:.0%}, AUC {n3(odr['auc'])}", f"Dự đoán điểm 3–9 rồi cắt ở ≈ {S['ord_cut']:.2f}; hơn SVM với p = {S['ord_vs_svm']['p']:.2f}".replace(".", ","), "Tận dụng thứ tự: điểm 6 gần lớp dương hơn điểm 4"], 16)

# 8 Ma trận nhầm lẫn
cb, cs = cms[BASE], cms[SVM]
sl = slide("Ma trận nhầm lẫn: đổi bỏ sót lấy báo nhầm", "Quyết định theo chi phí thực tế.")
pic(sl, "confusion_matrices.png", .6, 1.4, w=12.1)
bullets(sl, .6, 5.4, 12.1, 1.8, [f"Bỏ sót rượu ngon: {cb[1][0]} → {cs[1][0]}   |   Báo nhầm: {cb[0][1]} → {cs[0][1]}",
                                 "Precision ~0,4: sàng lọc tốt, nhưng chưa thay được hội đồng nếm"], 18)

# 9 Lỗi theo điểm
sl = slide("84% báo nhầm là rượu điểm 6 – sát ranh giới", "Nhãn 6/7 do con người đặt; điểm cảm quan có nhiễu.")
pic(sl, "pred_by_quality.png", .6, 1.5, h=4.6)
card(sl, 8.4, 1.6, 4.3, 4.8)
bullets(sl, 8.6, 1.8, 3.9, 4.5, ["Tỉ lệ đoán 'cao' tăng theo điểm thật: 9% (5) → 42% (6) → 80% (7)", "Bỏ sót: rượu ngon có độ cồn thấp", "Báo nhầm: rượu thường có độ cồn cao", "Mô hình phụ thuộc quá nhiều vào độ cồn"], 16)

# 10 Ngưỡng
sl = slide("Ngưỡng quan trọng ngang chọn mô hình", "Ngưỡng chọn trên dữ liệu đánh giá – lạc quan.")
pic(sl, "threshold.png", .6, 1.5, h=4.4)
table(sl, [["Mô hình", "F1 @0,5", "F1 tối ưu"]] + [[short(r["Mô hình"]), n3(r["F1 @0,5"]), n3(r["F1 tốt nhất"])] for _, r in thr_all.iterrows()], 7.6, 1.6, [2.8, 1.25, 1.25], 14)
tb(sl, 7.6, 5.0, 5.2, 1.5, "Sau khi chỉnh ngưỡng, các mô hình mạnh chỉ còn chênh 0,01–0,04", 16, False, INK)

# 11 Bản ghi trùng
sl = slide("Bản ghi trùng thổi phồng mô hình cây", "Rò rỉ train/test: cây nhớ bản sao.")
cd = CategoryChartData(); ms = ["Cơ sở: k-NN thô", "SVM RBF (C=10, cân bằng lớp)", "Random Forest (cân bằng lớp)", "Gradient Boosting (HistGB)"]
cd.categories = ["k-NN thô", "SVM RBF", "Random Forest", "Gradient Boosting"]
for tag, lab in [("Gốc", "Có trùng"), ("Đã", "Đã loại trùng")]:
    cd.add_series(lab, [round(ded[(ded["Dữ liệu"].str.startswith(tag)) & (ded["Mô hình"] == m)]["F1"].iloc[0], 3) for m in ms])
ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.6), Inches(1.4), Inches(7.8), Inches(5.6), cd).chart
ch.has_legend = True; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(14); pl = ch.plots[0]; pl.has_data_labels = True
pl.data_labels.font.size = Pt(12); pl.data_labels.number_format = "0.000"; pl.data_labels.number_format_is_linked = False; pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.has_major_gridlines = False; ch.value_axis.maximum_scale = .8; ch.value_axis.minimum_scale = 0; ch.category_axis.tick_labels.font.size = Pt(13); ch.value_axis.tick_labels.font.size = Pt(12)
for i, c in enumerate([SOFT, WINE]): pl.series[i].format.fill.solid(); pl.series[i].format.fill.fore_color.rgb = c
card(sl, 8.8, 1.6, 3.9, 4.8)
bullets(sl, 9.0, 1.8, 3.5, 4.5, [f"{S['n_dup']} bản ghi trùng (18%)", "Random Forest: F1 0,665 → 0,450", "SVM: chỉ 0,586 → 0,539", "Không loại trùng → chọn sai mô hình"], 16)

# 12 Hạn chế
sl = slide("Điều chưa giải quyết được", "Nói rõ giới hạn.")
for i, (h, t) in enumerate([("Nhãn chủ quan", "Ranh giới 6/7 có nhiễu, tạo trần hiệu năng"), ("Thiếu thông tin", "Không có giống nho, năm, hương vị; chỉ rượu Vinho Verde 2004–2007"),
                            ("Precision thấp (~0,4)", "Chỉ dùng làm công cụ sàng lọc, chưa thay được hội đồng nếm"), ("Đánh giá còn lạc quan", "Ngưỡng chọn trên dữ liệu ngoài-fold, chưa có tập kiểm tra độc lập")]):
    x, y = .6 + (i % 2) * 6.2, 1.6 + (i // 2) * 2.7
    card(sl, x, y, 5.9, 2.4); tb(sl, x + .25, y + .2, 5.4, .6, h, 20, True, WINE); tb(sl, x + .25, y + .85, 5.4, 1.5, t, 16, False, INK)

# 13 Demo
sl = slide("Demo: thao tác trực tiếp với dữ liệu", "streamlit run app.py")
bullets(sl, .6, 1.6, 7.2, 4.8, ["Chọn mẫu có sẵn hoặc lấy ngẫu nhiên một chai ngon", "Kéo 11 thanh chỉ tiêu hóa lý, đổi loại rượu", "So sánh k-NN cơ sở và SVM cải tiến", "Kéo ngưỡng để thấy đánh đổi bỏ sót / báo nhầm", "Đường xác suất theo độ cồn; vị trí trên PCA", "Tải CSV để dự đoán hàng loạt"], 19)
card(sl, 8.2, 1.9, 4.5, 3.0, WINE); tb(sl, 8.4, 2.4, 4.1, 2, ["streamlit run app.py", "→ localhost:8501"], 22, True, WHITE, PP_ALIGN.CENTER)

# 14 Kết luận
sl = prs.slides.add_slide(blank); bg(sl, WINE)
tb(sl, .8, .6, 11.5, 1, "Kết luận", 38, True, WHITE)
for i, (a, c) in enumerate([(f"{n3(b['f1'])} → {n3(s['f1'])}", "F1: chuẩn hóa + cân bằng lớp + SVM RBF"), (f"Thứ bậc: {n3(row('Hồi quy thứ bậc: SVR điểm → ngưỡng')['f1'])}", "SVR trên điểm: F1 cao nhất"), ("Loại trùng", "bắt buộc; lỗi dồn ở ranh giới 6/7")]):
    card(sl, .8 + i * 4.0, 2.0, 3.7, 3.0, hexc("74323F")); tb(sl, .95 + i * 4.0, 2.3, 3.4, 1, a, 28, True, WHITE, PP_ALIGN.CENTER); tb(sl, 1.0 + i * 4.0, 3.5, 3.3, 1.4, c, 17, False, SOFT, PP_ALIGN.CENTER)
tb(sl, .8, 5.6, 11.5, 1, "Phù hợp làm công cụ sàng lọc hỗ trợ, chưa thay thế đánh giá cảm quan.", 18, False, SOFT)

# 15 Phân công
sl = slide("Phân công và tài liệu tham khảo", "Điền tên thành viên.")
table(sl, [["Thành viên", "Nhiệm vụ", "Đóng góp"]] + [[m["name"], m["short"], m["share"]] for m in MEMBERS], .6, 1.6, [2.6, 3.4, 1.2], 16)
card(sl, 8.2, 1.6, 4.5, 4.8); tb(sl, 8.4, 1.75, 4.1, 4.5, ["[1] Cortez et al. 2009. Wine Quality, UCI: archive.ics.uci.edu/dataset/186", "[2] Cortez et al. 2009. Decision Support Systems 47(4)", "[3] Pedregosa et al. 2011. scikit-learn, JMLR 12", "[4] Nadeau & Bengio 2003. Machine Learning 52"], 13)
prs.save("Slide_NhanDangMau.pptx"); print("ok", len(prs.slides))
