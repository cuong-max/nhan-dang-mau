"""Sinh PhanCong_TuDanhGia.docx từ members.py."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from members import MEMBERS, SCHOOL, ADVISOR

d = Document()
sec = d.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(29.7), Cm(21)
sec.left_margin = sec.right_margin = Cm(2)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
p = d.add_paragraph(); p.add_run(SCHOOL).bold = True; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_heading("BẢNG PHÂN CÔNG VÀ TỰ ĐÁNH GIÁ ĐÓNG GÓP", 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph("Đề tài: Nhận dạng rượu vang chất lượng cao từ chỉ tiêu hóa lý (Wine Quality – UCI)")
d.add_paragraph("Học phần: Nhận dạng mẫu – Bài tập nhóm cuối kỳ")
d.add_paragraph(f"Giảng viên hướng dẫn: {ADVISOR}")

head = ["TT", "Họ tên", "MSSV", "Nhiệm vụ", "Sản phẩm phụ trách", "Tỉ lệ đóng góp", "Tự đánh giá (/10)"]
widths = [1, 3.6, 3, 7.2, 6.6, 2.2, 2.1]
t = d.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
for i, h in enumerate(head):
    c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(11)
for k, m in enumerate(MEMBERS, 1):
    cells = t.add_row().cells
    for c, v in zip(cells, [str(k), m["name"], m["mssv"], m["task"], m["products"], m["share"], ""]):
        c.text = ""; c.paragraphs[0].add_run(v).font.size = Pt(11)
cells = t.add_row().cells
cells[0].merge(cells[4]).text = "Tổng"; cells[5].text = "100%"
for row in t.rows:
    for c, w in zip(row.cells, widths): c.width = Cm(w)

d.add_paragraph()
d.add_paragraph("Tiêu chí tự đánh giá: mức độ hoàn thành nhiệm vụ, chất lượng sản phẩm, đúng hạn, tham gia thảo luận và hỗ trợ thành viên còn lại.")
d.save("PhanCong_TuDanhGia.docx"); print("ok")
