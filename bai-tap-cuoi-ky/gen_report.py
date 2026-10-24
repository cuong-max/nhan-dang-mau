"""Sinh BaoCao_NhanDangMau.docx từ results/ và figures/ (chạy sau run_experiments.py và robust_dedup.py)."""
from pathlib import Path
import pandas as pd, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from members import MEMBERS

R, F = Path("results"), Path("figures")
cmp_ = pd.read_csv(R / "model_comparison.csv"); sig = pd.read_csv(R / "significance_vs_baseline.csv").set_index("Mô hình")
S = json.load(open(R / "summary.json")); cms = json.load(open(R / "confusion_matrices.json"))
thr = pd.read_csv(R / "threshold_tradeoff.csv"); thr_all = pd.read_csv(R / "threshold_all_models.csv")
byq = pd.read_csv(R / "pred_rate_by_quality.csv"); bt = json.load(open(R / "by_wine_type.json"))
prof = pd.read_csv(R / "error_profile.csv", index_col=0); ded = pd.read_csv(R / "dedup_check.csv")
imp = pd.read_csv(R / "rf_importance.csv", index_col=0).iloc[:, 0]; qc = pd.read_csv(R / "quality_counts.csv")
row = lambda n: cmp_[cmp_["Mô hình"] == n].iloc[0]
BASE, SVM, ORD = "Cơ sở: k-NN (k=5, không chuẩn hóa)", "SVM RBF (cân bằng lớp, tuned)", "Hồi quy thứ bậc: SVR điểm → ngưỡng"
b, s = row(BASE), row(SVM)
pct = lambda v: f"{v*100:.1f}%".replace(".", ",")
zs = lambda v: f"{v:+.2f}".replace(".", ",").replace("-", "−")
num = lambda v, k=3: f"{v:.{k}f}".replace(".", ",")

d = Document()
sec = d.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.2)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.25
for n, sz in (("Heading 1", 15), ("Heading 2", 13)):
    h = d.styles[n]; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor(0x5A, 0x1E, 0x2B)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"): h.element.rPr.rFonts.set(qn(a), "Times New Roman")

def P(t, bold=False, italic=False, align=None, style=None, size=None):
    p = d.add_paragraph(style=style); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    return p
B = lambda t: P(t, style="List Bullet")
H = lambda t, l=1: d.add_heading(t, l)
C = WD_ALIGN_PARAGRAPH.CENTER
def shade(cell, color):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); cell._tc.get_or_add_tcPr().append(sh)
def table(rows, header, widths, hl=()):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10); shade(c, "EAD7DB")
    for ri, rw in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(rw):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(10); cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            if ri in hl: r.bold = True; shade(cells[i], "E2F0D9")
    for rw in t.rows:
        for i, w in enumerate(widths): rw.cells[i].width = Cm(w)
    return t
def cap(t): P(t, italic=True, align=C, size=11)
def fig(name, caption, w=15.5):
    d.add_picture(str(F / name), width=Cm(w)); d.paragraphs[-1].alignment = C; cap(caption)

# ---------- Bìa ----------
P("TRƯỜNG ........................................", True, align=C); P("KHOA ........................................", True, align=C)
for _ in range(4): d.add_paragraph()
P("BÁO CÁO BÀI TẬP NHÓM CUỐI KỲ", True, align=C, size=18); P("Học phần: Nhận dạng mẫu", align=C); d.add_paragraph()
P("NHẬN DẠNG RƯỢU VANG CHẤT LƯỢNG CAO TỪ CHỈ TIÊU HÓA LÝ", True, align=C, size=16)
P("Bộ dữ liệu Wine Quality (UCI) – so sánh phương án cơ sở với các phương án cải tiến", italic=True, align=C)
for _ in range(5): d.add_paragraph()
P("Nhóm thực hiện:", True, align=C)
for m in MEMBERS: P(f"{m['name']} – {m['mssv']}", align=C)
P("Giảng viên hướng dẫn: ........................................", align=C); P("Năm học 2026", align=C)
d.add_page_break()

# ---------- Tóm tắt ----------
H("Tóm tắt")
P(f"Báo cáo xây dựng bộ nhận dạng rượu vang chất lượng cao (điểm cảm quan ≥ 7) từ 11 chỉ tiêu hóa lý và loại rượu (đỏ/trắng), dùng bộ dữ liệu Wine Quality "
  f"của Cortez và cộng sự (UCI). Sau khi phát hiện và loại {S['n_dup']} bản ghi trùng lặp, dữ liệu còn {S['n']} mẫu, trong đó chỉ {pct(S['pos']/S['n'])} là lớp dương, nên "
  f"chúng tôi lấy F1 của lớp chất lượng cao và AUC làm thước đo chính thay vì accuracy. Phương án cơ sở là k-NN trên đặc trưng thô (F1 = {num(b['f1'])}); "
  f"trong các bộ phân lớp trực tiếp, SVM nhân RBF có chuẩn hóa và trọng số cân bằng lớp tốt nhất (F1 = {num(s['f1'])}, recall = {num(s['rec'])}, AUC = {num(s['auc'])}); "
  f"phương án thay thế tận dụng thứ tự của nhãn – hồi quy SVR trên điểm rồi đặt ngưỡng – đạt F1 cao nhất {num(row(ORD)['f1'])} (precision {num(row(ORD)['prec'])}). "
  f"Mọi cải thiện so với cơ sở đều có ý nghĩa thống kê (p < 0,001). Thử chọn đặc trưng theo gợi ý của UCI cho thấy nên giữ đủ 12 chỉ tiêu. Phân tích lỗi cho thấy (i) phần lớn báo nhầm đến từ rượu điểm 6 – sát ranh giới nhãn chủ quan 6/7, "
  f"(ii) bản ghi trùng làm thổi phồng kết quả của Random Forest/Gradient Boosting khoảng 0,2 điểm F1, (iii) ngưỡng quyết định ảnh hưởng tới F1 không kém việc chọn mô hình. "
  f"Nhóm xây dựng demo Streamlit cho phép thao tác trực tiếp với dữ liệu.")

# ---------- 1 ----------
H("1. Bài toán và mục tiêu")
P("Đánh giá chất lượng rượu vang hiện dựa vào hội đồng nếm thử – tốn kém, chậm và mang tính chủ quan. Ngược lại, các chỉ tiêu hóa lý (độ cồn, độ axit, đường dư, SO₂…) "
  "được đo tự động trong quá trình chứng nhận. Câu hỏi nhận dạng mẫu: chỉ từ vectơ 12 chiều này, có thể nhận ra rượu được hội đồng chấm điểm cao (≥ 7/10) hay không?")
P("Bộ dữ liệu gốc cho điểm 0–10 (bài toán hồi quy/thứ bậc). Nhóm chuyển thành bài toán phân lớp nhị phân 'chất lượng cao' và 'thường' vì: (a) đây là quyết định thực tế "
  "(chọn lô rượu cao cấp), (b) các điểm 3, 4, 8, 9 quá hiếm để học từng lớp riêng, (c) cho phép phân tích lỗi theo ma trận nhầm lẫn rõ ràng.")
P("Mục tiêu cụ thể:", bold=True)
B("Xây dựng phương án cơ sở và nhiều phương án cải tiến/thay thế, đánh giá công bằng trong cùng một quy trình kiểm định chéo.")
B("Xử lý mất cân bằng lớp (19% lớp dương) và kiểm định thống kê mức cải thiện.")
B("Phân tích lỗi: mẫu nào bị nhầm, vì sao; ảnh hưởng của bản ghi trùng, của ngưỡng và loại rượu; nêu hạn chế.")
B("Xây dựng demo tương tác.")

# ---------- 2 ----------
H("2. Dữ liệu")
H("2.1 Nguồn gốc", 2)
P("Wine Quality (Cortez, Cerdeira, Almeida, Matos & Reis, 2009), UCI Machine Learning Repository, https://archive.ics.uci.edu/dataset/186/wine+quality (doi: 10.24432/C56S3T). Dữ liệu thật thu thập 5/2004 – 2/2007 từ rượu "
  "Vinho Verde (Bắc Bồ Đào Nha) qua cơ quan chứng nhận CVRVV; mỗi mẫu được ít nhất 3 chuyên gia nếm mù và lấy trung vị điểm. Gồm 2 tệp: rượu đỏ (1.599) và trắng (4.898). "
  "Nhóm dùng bản tải về từ trang UCI (winequality-red.csv, winequality-white.csv, winequality.names), lưu trong thư mục data/; đã đối chiếu mã SHA-1 với bản trên máy chủ UCI. "
  "Lưu ý: trang UCI ghi 4.898 mẫu – đó là số mẫu rượu trắng; bài này dùng cả hai tệp (6.497 mẫu) và thêm biến is_red tương ứng với biến 'color' trên UCI.")
P("Mô tả trên UCI nêu ba điểm định hướng cách tiếp cận của nhóm: (1) các lớp có thứ tự và không cân bằng (rượu bình thường nhiều hơn hẳn rượu rất tốt/rất kém); "
  "(2) không chắc mọi biến đầu vào đều liên quan, nên đáng thử chọn đặc trưng; (3) chỉ có biến hóa lý và cảm quan, không có giống nho, nhãn hiệu, giá bán. "
  "Nhóm xử lý (1) bằng trọng số cân bằng lớp, chỉnh ngưỡng và một phương án hồi quy thứ bậc; (2) bằng một phương án chọn đặc trưng trong nested CV; (3) được nêu ở phần hạn chế.")
H("2.2 Mô tả và tiền xử lý", 2)
table([["Số mẫu gốc", f"{S['n_raw']} (đỏ 1.599 + trắng 4.898)"], ["Bản ghi trùng hoàn toàn", f"{S['n_dup']} ({pct(S['n_dup']/S['n_raw'])}) – bị loại"],
       ["Số mẫu sử dụng", f"{S['n']}"], ["Đặc trưng", "11 chỉ tiêu hóa lý + is_red (1 = đỏ)"],
       ["Nhãn", f"Chất lượng cao (điểm ≥ 7): {S['pos']} mẫu ({pct(S['pos']/S['n'])}); thường: {S['n']-S['pos']}"], ["Giá trị thiếu", "Không có"]],
      ["Thuộc tính", "Giá trị"], [4.5, 11.5])
cap("Bảng 1. Tóm tắt dữ liệu.")
P(f"Bản ghi trùng (cùng mọi chỉ tiêu và cùng điểm) chiếm {pct(S['n_dup']/S['n_raw'])}. Nếu giữ lại, cùng một mẫu có thể nằm cả ở tập huấn luyện lẫn tập kiểm tra, "
  "làm kết quả kiểm định chéo lạc quan giả tạo (rò rỉ dữ liệu). Mục 5.4 định lượng ảnh hưởng này.")
H("2.3 Khảo sát dữ liệu", 2)
fig("quality_dist.png", "Hình 1. Phân bố điểm chất lượng (sau loại trùng); màu đậm là lớp 'chất lượng cao'.", 10)
P("Phân bố điểm tập trung ở 5–6; điểm 3 và 9 cực hiếm. Ngưỡng 7 tạo lớp dương thiểu số, và hai điểm 6 và 7 – vốn chỉ cách nhau một bậc chủ quan – nằm ở hai phía ranh giới.")
fig("eda_hist.png", "Hình 2. Phân bố một số chỉ tiêu theo lớp (đã chuẩn hóa mật độ).")
hc = "; ".join(f"{k} (r = {num(v, 2)})" for k, v in S["high_corr"].items())
P(f"Độ cồn (alcohol) là chỉ tiêu phân biệt mạnh nhất, kế đến là tỉ trọng (density) – vốn tương quan âm mạnh với độ cồn. Các cặp tương quan |r| > 0,5: {hc}. "
  f"Hai lớp chồng lấp đáng kể trên mọi chỉ tiêu đơn lẻ, và trên mặt phẳng PCA (Hình 3) cũng không tách rõ: đây là bài toán khó hơn nhiều so với các bộ dữ liệu y sinh kinh điển. "
  f"Cần {S['pc95']}/12 thành phần chính để giữ 95% phương sai, tức thông tin trải đều, ít dư thừa.")
fig("pca.png", "Hình 3. Trái: chiếu PCA 2 chiều. Phải: phương sai tích lũy theo số thành phần chính.")

# ---------- 3 ----------
H("3. Phương pháp")
H("3.1 Phương án cơ sở", 2)
B("k-NN (k = 5) trên đặc trưng thô, không chuẩn hóa – phương án cơ sở chính (dùng cho kiểm định).")
B("Naive Bayes Gaussian trên đặc trưng thô – bộ phân loại Bayes đơn giản với giả thiết độc lập.")
B("Tham chiếu 'luôn đoán lớp đa số' – cho thấy vì sao accuracy gây hiểu nhầm (accuracy ≈ 81% nhưng không phát hiện được rượu ngon nào).")
H("3.2 Phương án cải tiến / thay thế", 2)
B("Chuẩn hóa z-score + k-NN (k = 15): sửa lỗi thang đo (total sulfur dioxide ~ 10² trong khi density ~ 1).")
B("LDA, QDA (reg_param = 0,1): bộ phân loại Bayes với phân phối Gauss đa biến, có xét tương quan.")
B("Hồi quy logistic L2 với trọng số cân bằng lớp, C ∈ {0,01; 0,1; 1; 10}.")
B(f"SVM nhân RBF với trọng số cân bằng lớp, C ∈ {{1; 10}}, γ ∈ {{scale; 0,2}} (tham số được chọn: C = {S['svm_params']['C']}, γ = {S['svm_params']['gamma']}); biến thể PCA (8 thành phần) + SVM.")
B("Random Forest 300 cây (balanced_subsample) và Histogram Gradient Boosting – mô hình tập hợp dựa trên cây.")
B("Chọn đặc trưng kiểu lọc (thống kê ANOVA F) + SVM RBF: số đặc trưng k ∈ {3, 5, 7, 9, 12} được chọn trong nested CV, cùng với C.")
B("Hồi quy thứ bậc: SVR (nhân RBF) dự đoán điểm 3–9 thay vì nhãn nhị phân, rồi gán 'chất lượng cao' nếu điểm dự đoán ≥ ngưỡng; ngưỡng được chọn bằng CV nội bộ trên tập huấn luyện. "
  "Cách này tận dụng thứ tự của nhãn (rượu điểm 6 'gần' lớp dương hơn điểm 4).")
B("Điều chỉnh ngưỡng quyết định trên xác suất (Platt scaling cho SVM) nhằm tối đa F1.")
H("3.3 Quy trình đánh giá", 2)
B("Kiểm định chéo phân tầng 5-fold lặp 3 lần (15 fold); chuẩn hóa/PCA nằm trong Pipeline, chỉ fit trên tập huấn luyện của từng fold.")
B("Mô hình có siêu tham số được chọn bằng GridSearchCV 3-fold (tiêu chí AUC) bên trong mỗi fold huấn luyện (nested CV).")
B("Chỉ số: F1, recall, precision của lớp chất lượng cao, balanced accuracy, ROC-AUC, average precision; accuracy chỉ để tham khảo.")
B("Kiểm định t ghép cặp với hiệu chỉnh phương sai Nadeau–Bengio trên F1 của 15 fold, so với k-NN thô.")
B("Phân tích lỗi trên dự đoán ngoài-fold (5-fold) của toàn bộ mẫu; seed cố định = 42.")

# ---------- 4 ----------
H("4. Kết quả")
rows = []
for _, r in cmp_.iterrows():
    n = r["Mô hình"]; pv = sig.loc[n, "p"] if n in sig.index else None
    rows.append([n, f"{num(r['f1'])} ± {num(r['f1_sd'])}", num(r["rec"]), num(r["prec"]), num(r["bacc"]), num(r["auc"]), num(r["acc"]),
                 "–" if pv is None else f"{r['f1']-b['f1']:+.3f}".replace(".", ",").replace("-", "−") + ("; p<0,001" if pv < .001 else f"; p={num(pv)}")])
table(rows, ["Mô hình", "F1", "Recall", "Precision", "Bal. acc", "AUC", "Acc", "ΔF1 so với cơ sở"], [4.6, 1.9, 1.3, 1.5, 1.3, 1.2, 1.2, 2.6], hl=(0,))
cap("Bảng 2. Kết quả trung bình 15 fold, xếp theo F1 giảm dần (ngưỡng mặc định). p theo t-test hiệu chỉnh Nadeau–Bengio.")
fig("cv_boxplot.png", "Hình 4. Phân bố F1 trên 15 fold.", 13.5)
H("4.1 Nhận xét", 2)
dm = row("Tham chiếu: luôn đoán lớp đa số")
B(f"Accuracy gây hiểu nhầm: bộ 'luôn đoán thường' đạt accuracy {pct(dm['acc'])}, cao hơn cả SVM ({pct(s['acc'])}), nhưng F1 = 0. Phương án cơ sở k-NN thô chỉ bắt được "
  f"{pct(b['rec'])} rượu ngon (F1 = {num(b['f1'])}).")
B(f"Chuẩn hóa giúp k-NN tăng F1 từ {num(b['f1'])} lên {num(row('k-NN + chuẩn hóa')['f1'])} và AUC từ {num(b['auc'])} lên {num(row('k-NN + chuẩn hóa')['auc'])}: khoảng cách Euclid "
  "trên dữ liệu thô bị chi phối bởi SO₂ tự do/tổng (giá trị hàng chục–hàng trăm), trong khi độ cồn – đặc trưng quan trọng nhất – chỉ dao động 8–14.")
B(f"Trọng số cân bằng lớp là yếu tố quyết định F1: SVM và hồi quy logistic (có cân bằng lớp) đạt recall {num(row('Hồi quy logistic (cân bằng lớp, tuned)')['rec'],2)}–{num(s['rec'],2)} "
  "nhưng precision chỉ ~0,4. QDA, nhờ ước lượng hiệp phương sai riêng cho lớp thiểu số, cũng có recall cao (0,69) mà không cần trọng số. Ba mô hình SVM, QDA, hồi quy logistic chênh nhau nhỏ hơn độ lệch chuẩn giữa các fold nên coi như tương đương; SVM RBF nhỉnh nhất trong nhóm phân lớp trực tiếp, nên được dùng cho phân tích lỗi ở Mục 5 và cho demo.")
rf, hg = row("Random Forest (cân bằng lớp)"), row("Gradient Boosting (HistGB)")
B(f"Random Forest có AUC cao nhất ({num(rf['auc'])}) nhưng F1 ở ngưỡng 0,5 thấp ({num(rf['f1'])}) do recall chỉ {num(rf['rec'],2)}: mô hình xếp hạng tốt nhưng ngưỡng mặc định không phù hợp "
  "với lớp thiểu số. Mục 5.3 cho thấy khi chỉnh ngưỡng, khoảng cách giữa các mô hình thu hẹp đáng kể.")
fsr, odr = row("Chọn đặc trưng (ANOVA, k tuned) + SVM"), row("Hồi quy thứ bậc: SVR điểm → ngưỡng")
B(f"Chọn đặc trưng: khi chọn trên toàn bộ dữ liệu, nested CV giữ k = {S['fs_k']} đặc trưng ({', '.join(S['fs_kept'])}). F1 = {num(fsr['f1'])}, AUC = {num(fsr['auc'])} so với SVM đầy đủ "
  f"{num(s['f1'])} / {num(s['auc'])}. " + ("Bỏ bớt đặc trưng không giúp tăng chất lượng: các chỉ tiêu có F-score thấp vẫn mang thông tin bổ sung cho mô hình phi tuyến, khớp với kết quả ở Mục 5.5." if fsr['f1'] <= s['f1'] + .005 else "Bỏ bớt đặc trưng nhiễu giúp mô hình tốt hơn đôi chút."))
B(f"Hồi quy thứ bậc (SVR): F1 = {num(odr['f1'])}, recall = {num(odr['rec'])}, precision = {num(odr['prec'])}, AUC = {num(odr['auc'])} (ngưỡng điểm trung bình {num(S['ord_cut'], 2)}); hơn SVM {num(S['ord_vs_svm']['dF1'])} điểm F1 với p = {num(S['ord_vs_svm']['p'])} – ý nghĩa thống kê ở mức vừa phải. "
  + ("Tận dụng thứ tự của nhãn cho kết quả ngang hoặc tốt hơn các bộ phân loại nhị phân tốt nhất, một hướng đáng chú ý cho bài toán này." if odr['f1'] >= s['f1'] - .01 else "Kết quả thấp hơn SVM phân lớp trực tiếp: việc học cả các khác biệt 3/4/5 không giúp ích nhiều cho ranh giới 6/7."))
B("LDA không cải thiện F1 (biên tuyến tính, không bù mất cân bằng); PCA + SVM kém SVM đầy đủ một chút – thông tin không tập trung vào vài thành phần đầu (khớp với Hình 3).")
fig("roc_pr.png", "Hình 5. Đường ROC và Precision–Recall (dự đoán ngoài-fold).")

# ---------- 5 ----------
H("5. Phân tích lỗi")
H("5.1 Ma trận nhầm lẫn", 2)
fig("confusion_matrices.png", "Hình 6. Ma trận nhầm lẫn (dự đoán ngoài-fold).", 16)
cb, cl, cs = cms[BASE], cms["Hồi quy logistic (cân bằng lớp, tuned)"], cms[SVM]
P(f"Cơ sở bỏ sót {cb[1][0]}/{S['pos']} rượu ngon. SVM giảm bỏ sót xuống {cs[1][0]} nhưng đổi lại {cs[0][1]} báo nhầm (so với {cb[0][1]}). Với bài toán chọn lô cao cấp, "
  "đây là đánh đổi phải được quyết định theo chi phí thực tế: nếu báo nhầm đắt (bán rượu thường với giá cao), cần nâng ngưỡng.")
H("5.2 Lỗi tập trung ở ranh giới điểm 6/7", 2)
fig("pred_by_quality.png", "Hình 7. Tỉ lệ SVM dự đoán 'chất lượng cao' theo điểm thực tế.", 11)
q6 = byq[byq["quality"] == 6].iloc[0]; q5 = byq[byq["quality"] == 5].iloc[0]; q7 = byq[byq["quality"] == 7].iloc[0]; q8 = byq[byq["quality"] == 8].iloc[0]
fp6 = q6["n"] * q6.iloc[2]
P(f"Mô hình dự đoán 'cao' cho {pct(q6.iloc[2])} rượu điểm 6 nhưng chỉ {pct(q5.iloc[2])} rượu điểm 5; với điểm 7 và 8 lần lượt là {pct(q7.iloc[2])} và {pct(q8.iloc[2])}. "
  f"Như vậy khoảng {int(round(fp6))}/{cs[0][1]} báo nhầm ({pct(fp6/cs[0][1])}) là rượu điểm 6. Từ điểm 5 trở lên, tỉ lệ dự đoán 'cao' tăng dần theo điểm thật, cho thấy mô hình nắm được thứ bậc chất lượng "
  "(điểm 3–4 quá ít mẫu để kết luận); lỗi chủ yếu do ranh giới 6/7 là quy ước của con người, trong khi điểm cảm quan vốn có nhiễu.")
H("5.3 Đặc điểm mẫu bị nhầm và ảnh hưởng của ngưỡng", 2)
feat_show = ["alcohol", "density", "volatile acidity", "chlorides", "residual sugar"]
table([[f] + [zs(prof.loc[f, c]) for c in ["TP", "FN", "FP", "TN"]] for f in feat_show], ["Chỉ tiêu (z-score TB)", "TP", "FN (bỏ sót)", "FP (báo nhầm)", "TN"], [5, 2.5, 2.8, 2.8, 2.5])
cap("Bảng 3. Giá trị chuẩn hóa trung bình theo nhóm kết quả (SVM).")
P(f"Rượu ngon bị bỏ sót có độ cồn dưới trung bình (z = {zs(prof.loc['alcohol','FN'])}) và tỉ trọng cao (z = {zs(prof.loc['density','FN'])}) – trái ngược với rượu ngon được nhận đúng "
  f"(alcohol z = {zs(prof.loc['alcohol','TP'])}). Ngược lại, rượu thường bị báo nhầm có độ cồn cao (z = {zs(prof.loc['alcohol','FP'])}). Mô hình dựa quá nhiều vào độ cồn; "
  "những yếu tố khác làm nên chất lượng (giống nho, hương vị, cân bằng) không có trong dữ liệu.")
fig("threshold.png", "Hình 8. F1, precision, recall theo ngưỡng (SVM, xác suất Platt).", 11)
table([[n, num(r["F1 @0,5"]), num(r["F1 tốt nhất"]), num(r["Ngưỡng"], 2), num(r["AUC"])] for n, r in thr_all.set_index("Mô hình").iterrows()],
      ["Mô hình", "F1 @ 0,5", "F1 sau chỉnh ngưỡng", "Ngưỡng", "AUC"], [6.5, 2.2, 3, 2, 2])
cap("Bảng 4. F1 trước và sau khi chọn ngưỡng tối ưu (dự đoán ngoài-fold). Ngưỡng chọn trên chính dữ liệu đánh giá nên số liệu lạc quan.")
P("Khi mọi mô hình được chỉnh ngưỡng, Random Forest đạt F1 cao nhất và khoảng cách giữa các mô hình mạnh chỉ còn 0,01–0,04. Bài học: với dữ liệu mất cân bằng, "
  "chọn ngưỡng (hoặc trọng số lớp) quan trọng ngang chọn thuật toán; so sánh mô hình nên dựa trên AUC/AP hoặc sau khi chỉnh ngưỡng trên tập kiểm định riêng.")
H("5.4 Bản ghi trùng làm thổi phồng kết quả", 2)
dd = ded.copy(); dd["Dữ liệu"] = dd["Dữ liệu"].str.replace("Gốc", "Có trùng")
table([[r["Dữ liệu"], r["Mô hình"], num(r["F1"]), num(r["AUC"]), num(r["Recall"])] for _, r in dd.iterrows()], ["Dữ liệu", "Mô hình", "F1", "AUC", "Recall"], [4, 6, 2, 2, 2])
cap("Bảng 5. Cùng quy trình đánh giá, có và không có bản ghi trùng.")
g = lambda tag, m: dd[(dd["Dữ liệu"].str.startswith(tag)) & (dd["Mô hình"].str.startswith(m))].iloc[0]["F1"]
P(f"Với dữ liệu còn trùng, Random Forest đạt F1 {num(g('Có', 'Random'))} và Gradient Boosting {num(g('Có', 'Gradient'))}, có vẻ vượt xa SVM ({num(g('Có', 'SVM'))}). "
  f"Sau khi loại trùng, F1 của chúng rơi xuống {num(g('Đã', 'Random'))} và {num(g('Đã', 'Gradient'))}, trong khi SVM chỉ giảm nhẹ ({num(g('Đã', 'SVM'))}). "
  "Nguyên nhân: mô hình cây có thể 'nhớ' chính xác mẫu huấn luyện, nên khi bản sao nằm ở tập kiểm tra chúng được chấm đúng gần như chắc chắn. "
  "Đây là lỗi phương pháp thường gặp với bộ dữ liệu này; nhiều kết quả công bố trên Wine Quality không loại trùng nên có thể lạc quan.")
H("5.5 Theo loại rượu, số đặc trưng và lượng dữ liệu", 2)
P(f"F1 của SVM tương đương giữa rượu đỏ ({num(bt['Đỏ']['f1'])}) và trắng ({num(bt['Trắng']['f1'])}); AUC trên rượu đỏ cao hơn ({num(bt['Đỏ']['auc'])} so với {num(bt['Trắng']['auc'])}). "
  "Huấn luyện chung hai loại và thêm biến is_red không làm hại rượu đỏ, dù chỉ chiếm 26% mẫu.")
fig("sensitivity.png", "Hình 9. Trái: AUC của SVM theo số đặc trưng quan trọng nhất (RF). Phải: đường cong học (5 lần chia 80/20).")
P(f"Chỉ dùng độ cồn đã đạt AUC {num(S['k_auc'][0])}; đủ 12 đặc trưng đạt {num(S['k_auc'][-1])} – các chỉ tiêu còn lại có đóng góp nhưng mỗi cái nhỏ. "
  f"Đường cong học của SVM và Random Forest vẫn còn đi lên khi dùng 100% dữ liệu (AUC {num(S['lc_svm'][-1])} / {num(S['lc_rf'][-1])}), gợi ý thêm dữ liệu còn có ích; "
  f"k-NN thô kém hẳn ở mọi cỡ dữ liệu ({num(S['lc_knn'][-1])}).")

# ---------- 6 ----------
H("6. Hạn chế")
B("Nhãn là điểm cảm quan chủ quan (trung vị ≥ 3 chuyên gia); ranh giới 6/7 có nhiễu nên tồn tại trần hiệu năng mà không mô hình nào vượt qua được bằng các đặc trưng hiện có.")
B("Chỉ có chỉ tiêu hóa lý – không có giống nho, năm, giá bán, hương vị; vì lý do bảo mật, nhãn hiệu không được công bố. Kết quả chỉ áp dụng cho rượu Vinho Verde 2004–2007.")
B("Precision của mô hình tốt nhất chỉ ~0,4 ở ngưỡng mặc định: cứ 10 chai được đề xuất 'ngon' thì khoảng 6 chai thực ra là rượu thường – chưa đủ để thay thế hội đồng nếm.")
B("Ngưỡng tối ưu và phân tích lỗi tính trên dự đoán ngoài-fold của toàn bộ dữ liệu, không có tập kiểm tra độc lập hoàn toàn; ngưỡng tối ưu vì vậy hơi lạc quan.")
B("Lưới siêu tham số nhỏ để giới hạn thời gian chạy; phương án thứ bậc mới dừng ở SVR + ngưỡng, chưa thử mô hình thứ bậc chuyên dụng. Demo dùng SVM phân lớp vì cần xác suất.")
B("Demo dùng mô hình huấn luyện trên toàn bộ dữ liệu nên dự đoán cho mẫu có sẵn là trong-mẫu.")
H("Hướng phát triển", 2)
B("Mô hình thứ bậc chuyên dụng (ordinal logistic, cost-sensitive theo khoảng cách điểm); chọn ngưỡng trên tập kiểm định riêng theo hàm chi phí của nhà sản xuất.")
B("Theo gợi ý của UCI: dùng phát hiện ngoại lai (outlier/novelty detection) để tìm các rượu rất tốt hoặc rất kém – nhóm hiếm (điểm 3, 9).")
B("Hiệu chuẩn xác suất và giải thích (SHAP) để người dùng hiểu vì sao một mẫu được đề xuất.")
B("Bổ sung đặc trưng ngoài hóa lý (giống nho, phổ hương) và dữ liệu từ vùng/năm khác để kiểm chứng ngoài.")

# ---------- 7 ----------
H("7. Kết luận")
P(f"Trên bài toán nhận dạng rượu vang chất lượng cao, phương án cải tiến SVM RBF + chuẩn hóa + cân bằng lớp nâng F1 từ {num(b['f1'])} (k-NN thô) lên {num(s['f1'])} và AUC từ "
  f"{num(b['auc'])} lên {num(s['auc'])}, có ý nghĩa thống kê; hồi quy thứ bậc SVR, vốn tận dụng thứ tự của điểm chất lượng, còn đạt F1 {num(row(ORD)['f1'])}. "
  "Bốn yếu tố tạo ra cải thiện: chuẩn hóa đặc trưng, xử lý mất cân bằng lớp, lựa chọn ngưỡng và khai thác thứ bậc của nhãn; còn chọn đặc trưng thì không giúp ích. "
  "Phân tích lỗi chỉ ra rằng phần lớn sai số nằm ở ranh giới điểm 6/7 và ở những rượu ngon có độ cồn thấp; đồng thời việc loại bản ghi trùng là bắt buộc để đánh giá trung thực. "
  "Mô hình phù hợp làm công cụ sàng lọc hỗ trợ, chưa thể thay thế đánh giá cảm quan.")

# ---------- 8 ----------
H("8. Mã nguồn và hướng dẫn chạy")
table([["run_experiments.py", "Tải dữ liệu, loại trùng, chạy toàn bộ thí nghiệm, sinh results/, figures/, model_best.joblib"],
       ["robust_dedup.py", "So sánh kết quả có/không bản ghi trùng (Bảng 5)"], ["app.py", "Demo Streamlit"],
       ["gen_report.py, gen_slides.py", "Sinh báo cáo .docx và slide .pptx từ kết quả"], ["data/", "2 tệp CSV gốc từ UCI"], ["results/, figures/", "Bảng số liệu và hình"]],
      ["Tệp", "Chức năng"], [4.5, 11.5])
d.add_paragraph()
for c in ["pip install -r requirements.txt", "python run_experiments.py     # khoảng 10 phút; dùng ./data", "python run_experiments.py --data-dir ~/Downloads/wine+quality   # thư mục dữ liệu khác", "python robust_dedup.py", "streamlit run app.py          # mở http://localhost:8501"]:
    p = P(c); p.runs[0].font.name = "Courier New"; p.runs[0].font.size = Pt(10.5)

# ---------- 9 ----------
H("9. Phân công và tự đánh giá đóng góp")
table([[str(i + 1), f"{m['name']}\n{m['mssv']}", m["task"], m["share"], "...."] for i, m in enumerate(MEMBERS)],
      ["TT", "Họ tên / MSSV", "Nhiệm vụ", "Đóng góp", "Tự đánh giá (/10)"], [1, 3.5, 7.2, 2, 2.3])
P("Tổng đóng góp: 100%. Hai thành viên cùng thảo luận chọn phương án và thống nhất kết luận.", italic=True)

# ---------- TLTK ----------
H("Tài liệu tham khảo và nguồn")
for r in ["[1] Cortez, P., Cerdeira, A., Almeida, F., Matos, T., Reis, J. (2009). Wine Quality [Dataset]. UCI Machine Learning Repository. https://archive.ics.uci.edu/dataset/186/wine+quality, https://doi.org/10.24432/C56S3T",
          "[2] Cortez, P., Cerdeira, A., Almeida, F., Matos, T., Reis, J. (2009). Modeling wine preferences by data mining from physicochemical properties. Decision Support Systems 47(4), 547–553.",
          "[3] Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR 12, 2825–2830. https://scikit-learn.org (mã tham khảo: Pipeline, GridSearchCV, SVC, HistGradientBoostingClassifier).",
          "[4] Nadeau, C., Bengio, Y. (2003). Inference for the generalization error. Machine Learning 52(3), 239–281.",
          "[5] Cortes, C., Vapnik, V. (1995). Support-vector networks. Machine Learning 20(3), 273–297.",
          "[6] Breiman, L. (2001). Random Forests. Machine Learning 45(1), 5–32.",
          "[7] Duda, R. O., Hart, P. E., Stork, D. G. (2001). Pattern Classification, 2nd ed. Wiley.",
          "[8] Kaufman, S., Rosset, S., Perlich, C. (2012). Leakage in data mining: formulation, detection, and avoidance. ACM TKDD 6(4).",
          "[9] Streamlit documentation. https://docs.streamlit.io"]:
    P(r)
d.save("BaoCao_NhanDangMau.docx"); print("ok")
